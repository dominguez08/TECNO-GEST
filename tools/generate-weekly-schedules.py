from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / 'docs' / 'cronogramas'
BLUE, PALE = '173F73', 'EFF5FC'
DAYS = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes']

WEEKS = [
 ('Semana 1', 'Análisis y planificación', 'Requerimientos y planificación del sistema',
  [('Revisar necesidades del inventario', 'Definir estructura visual, navegación y pantallas necesarias'), ('Definir usuarios y casos de uso', 'Crear bocetos del inicio de sesión y panel principal'), ('Organizar módulos y prioridades', 'Diseñar la guía visual: colores, tipografía y componentes'), ('Planificar base de datos y flujos', 'Preparar prototipos de inventario, préstamos y mantenimiento'), ('Consolidar cronograma y evidencias', 'Ajustar prototipos y documentar decisiones de interfaz')],
  'Plan del proyecto, requerimientos y prototipos iniciales.'),
 ('Semana 2', 'Diseño del sistema', 'Diseño de base de datos, interfaz y módulos',
  [('Diseñar tablas, relaciones y restricciones', 'Construir estructura HTML de acceso y navegación'), ('Definir campos de equipos, usuarios y ubicaciones', 'Aplicar estilos CSS generales y diseño adaptable'), ('Preparar esquema de préstamos y mantenimiento', 'Diseñar las pantallas de panel e inventario'), ('Configurar conexión y estructura de carpetas', 'Diseñar formularios de equipos y ubicaciones'), ('Revisar diseño con el integrante', 'Integrar plantillas visuales y validar navegación')],
  'Esquema de datos, estructura del proyecto y diseño de interfaces.'),
 ('Semana 3', 'Desarrollo inicial', 'Primeros módulos funcionales',
  [('Implementar autenticación, sesiones y roles', 'Implementar vistas de inicio de sesión y panel'), ('Programar consultas del panel e inventario', 'Construir listado de inventario, búsqueda y filtros'), ('Crear operaciones para equipos y ubicaciones', 'Desarrollar formularios de registro y edición'), ('Validar datos en servidor y consultas preparadas', 'Mostrar mensajes, estados y datos de las fichas'), ('Probar flujo de registro de equipos', 'Corregir detalles visuales del inventario')],
  'Acceso al sistema, panel, inventario y ubicaciones operativos.'),
 ('Semana 4', 'Desarrollo e integración', 'Sistema integrado con préstamos y mantenimiento',
  [('Programar creación, devolución y reglas de préstamos', 'Crear vistas y formularios de préstamos'), ('Implementar reportes de fallas y mantenimiento', 'Diseñar listado y detalle de mantenimiento'), ('Aplicar transacciones y validaciones de disponibilidad', 'Integrar estados, etiquetas y acciones de devolución'), ('Conectar reportes con los datos reales', 'Adaptar las pantallas para dispositivos móviles'), ('Verificar integración entre módulos', 'Revisar consistencia visual de la aplicación')],
  'Préstamos, devoluciones, mantenimiento y vistas integradas.'),
 ('Semana 5', 'Pruebas correcciones y documentación', 'Sistema probado y avance de manuales',
  [('Realizar pruebas de roles, préstamos y mantenimiento', 'Probar formularios, navegación y visualización móvil'), ('Corregir validaciones y mensajes de error', 'Corregir estilos, espaciados y retroalimentación visual'), ('Verificar seguridad de sesiones y formularios', 'Preparar capturas reales para el manual de usuario'), ('Comprobar reportes y exportación CSV', 'Documentar procedimientos de las pantallas'), ('Registrar resultados de pruebas', 'Organizar evidencias visuales y revisar manual')],
  'Correcciones aplicadas, pruebas registradas y documentación en avance.'),
 ('Semana 6', 'Finalización y presentación', 'Proyecto y documentación final',
  [('Revisar instalación, base de datos y respaldos', 'Realizar revisión final de todas las pantallas'), ('Ejecutar pruebas finales de funcionamiento', 'Ajustar visualización final en computadora y teléfono'), ('Verificar permisos y protección de datos', 'Completar capturas y manual de usuario final'), ('Preparar demostración técnica del sistema', 'Preparar recorrido visual para la presentación'), ('Entregar código y documentación final', 'Entregar manuales y evidencias finales')],
  'Sistema InventIC terminado, documentación y presentación preparadas.')
]

def set_font(run, size=10, bold=False, color=None):
    run.font.name = 'Aptos'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos'); run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, color):
    tag = OxmlElement('w:shd'); tag.set(qn('w:fill'), color); cell._tc.get_or_add_tcPr().append(tag)

def add_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, item in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=''; shade(c, BLUE); r=c.paragraphs[0].add_run(item); set_font(r,9,True,'FFFFFF')
    for ri, row in enumerate(rows):
        cells=t.add_row().cells
        for i, item in enumerate(row):
            cells[i].text=''; p=cells[i].paragraphs[0]; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2); r=p.add_run(item); set_font(r,9.1)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2: shade(cells[i], PALE)
    for row in t.rows:
        for c,w in zip(row.cells,widths): c.width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)

def schedule_table(items, role):
    return [(day, task, detail, '5 h') for day, (task, detail) in zip(DAYS, items)] if role == 'backend' else [(day, detail, task, '5 h') for day, (task, detail) in zip(DAYS, items)]

def build(index, week, stage, product, tasks, result):
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    normal=d.styles['Normal']; normal.font.name='Aptos'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); normal.font.size=Pt(10)
    for st, size in [('Title',22),('Heading 1',14),('Heading 2',11)]:
        s=d.styles[st]; s.font.name='Aptos'; s._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0); s.font.bold=True
    foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=foot.add_run(f'InventIC | Cronograma {week} | 2026'); set_font(r,8,False,'666666')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('INVENTIC IEP SAN RAFAEL'); set_font(r,18,True,BLUE)
    p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f'Cronograma Individual {week}')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(stage); set_font(r,11)
    add_table(d,['Elemento','Información'], [('Nombre del proyecto','InventIC IEP San Rafael'),('Integrantes','Jesús Ernesto Domínguez Reyes - Front end\nJosé Baltazar Beltrán Gómez - Back end'),('Semana',week),('Etapa principal',stage),('Producto esperado',product),('Carga individual','25 horas por integrante')],[1.65,5.15])
    d.add_heading('Objetivo semanal',1); p=d.add_paragraph(f'Durante {week.lower()} se desarrollan las actividades de {stage.lower()}. El resultado esperado es: {result}')
    d.add_heading('Cronograma de Jesús Ernesto Domínguez Reyes',1); p=d.add_paragraph(); r=p.add_run('Rol: Front end. '); set_font(r,10,True); p.add_run('Desarrollo de interfaces, estilos, navegación y experiencia de usuario.')
    add_table(d,['Día','Actividad','Tarea','Horas'],schedule_table(tasks,'frontend'),[.85,1.8,3.2,.65])
    d.add_heading('Cronograma de José Baltazar Beltrán Gómez',1); p=d.add_paragraph(); r=p.add_run('Rol: Back end. '); set_font(r,10,True); p.add_run('Desarrollo de lógica, datos, seguridad, validaciones e integración.')
    add_table(d,['Día','Actividad','Tarea','Horas'],schedule_table(tasks,'backend'),[.85,1.8,3.2,.65])
    d.add_heading('Evidencia de cierre',1); d.add_paragraph(f'Producto para comprobar la semana: {result} Se deben conservar capturas de pantalla y el avance del código relacionado.')
    d.core_properties.title=f'Cronograma {week} InventIC'; d.core_properties.author='Jesús Ernesto Domínguez Reyes y José Baltazar Beltrán Gómez'
    d.save(OUT / f'{index:02d}_Cronograma_{week.replace(" ", "_")}.docx')

OUT.mkdir(parents=True, exist_ok=True)
for number, record in enumerate(WEEKS,1): build(number,*record)
print(OUT)
