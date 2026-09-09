from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'manuales'
OUT.mkdir(parents=True, exist_ok=True)
BLUE = '07356B'
PALE = 'EAF2FB'
GRAY = 'D9E1EA'

def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tc_pr.append(shd)

def borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for n in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{n}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), GRAY); b.append(e)
    tbl_pr.append(b)

def set_cell(cell, text, bold=False, color=None, size=9.2):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    r = p.add_run(str(text)); r.bold = bold; r.font.name = 'Aptos'; r._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos'); r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, data, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'; borders(t)
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, 'FFFFFF'); shade(t.rows[0].cells[i], BLUE)
    for ri, row in enumerate(data):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if ri % 2: shade(cells[i], 'F7FAFD')
    if widths:
        for row in t.rows:
            for cell, width in zip(row.cells, widths): cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

def base_doc(kind):
    d = Document(); sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.65)
    sec.left_margin, sec.right_margin = Inches(.72), Inches(.72)
    normal = d.styles['Normal']; normal.font.name = 'Aptos'; normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos'); normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.12
    for name, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 12), ('Heading 3', 10.8)]:
        s = d.styles[name]; s.font.name = 'Aptos Display' if name != 'Heading 3' else 'Aptos'; s._element.rPr.rFonts.set(qn('w:ascii'), s.font.name); s._element.rPr.rFonts.set(qn('w:hAnsi'), s.font.name); s.font.size = Pt(size); s.font.color.rgb = RGBColor(0,0,0); s.font.bold = True
        s.paragraph_format.space_before = Pt(12 if name != 'Title' else 0); s.paragraph_format.space_after = Pt(6)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('InventIC  |  ' + kind + '  |  Septiembre de 2026'); run.font.size = Pt(8); run.font.color.rgb = RGBColor(80,80,80)
    return d

def cover(doc, document_name, subtitle):
    for _ in range(5): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run('INVENTIC'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(30); r.font.color.rgb=RGBColor.from_string(BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Sistema de Inventario de Equipos'); r.font.size=Pt(15); r.font.color.rgb=RGBColor(30,30,30)
    doc.add_paragraph()
    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(document_name)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(subtitle); r.italic=True; r.font.size=Pt(12)
    for _ in range(5): doc.add_paragraph()
    table(doc, ['Dato', 'Información'], [
        ['Institución', 'IEP San Rafael'], ['Especialidad', 'Desarrollo de Software'], ['Proyecto', 'InventIC - TECNO-GEST'],
        ['Integrantes', '________________________________________'], ['Docente', '________________________________________'],
        ['Versión del documento', '1.0'], ['Fecha', '9 de septiembre de 2026'],
    ], [1.8, 4.9])
    doc.add_page_break()

def heading(doc, text, level=1): doc.add_heading(text, level=level)
def para(doc, text, bold_start=None):
    p=doc.add_paragraph()
    if bold_start and text.startswith(bold_start):
        p.add_run(bold_start).bold=True; p.add_run(text[len(bold_start):])
    else: p.add_run(text)
    return p
def bullets(doc, items):
    for x in items: doc.add_paragraph(x, style='List Bullet')
def steps(doc, items):
    for x in items: doc.add_paragraph(x, style='List Number')
def version_table(doc):
    table(doc, ['Versión', 'Fecha', 'Cambio', 'Responsable'], [['1.0', '09/09/2026', 'Primera versión basada en la versión actual del sistema.', 'Equipo de desarrollo']], [0.75,1.15,3.7,1.1])

def user_manual():
    d=base_doc('Manual de Usuario'); cover(d, 'Manual de Usuario', 'Guía para operar el sistema institucional de inventario')
    heading(d, 'Control de versiones'); version_table(d)
    heading(d, '1 Introducción')
    para(d, 'InventIC es una aplicación web para registrar y consultar equipos institucionales, controlar préstamos, reportar fallas y dar seguimiento al mantenimiento. Centraliza la información de los equipos por sede y ubicación para que el inventario se mantenga actualizado.')
    para(d, 'Este manual está dirigido a personal administrativo, técnico y docente. Explica las acciones disponibles de acuerdo con el perfil de acceso y utiliza los mismos nombres que aparecen en la interfaz.')
    heading(d, '2 Requisitos para utilizar el sistema')
    table(d, ['Elemento', 'Requisito'], [['Dispositivo','Computadora, laptop, tableta o teléfono con navegador moderno.'],['Navegador','Google Chrome, Microsoft Edge, Firefox o Safari actualizado.'],['Conexión','Red local o Internet según la instalación institucional.'],['Cuenta','Correo y contraseña asignados por el administrador.']], [1.4,5.2])
    heading(d, '3 Acceso al sistema')
    steps(d, ['Abra el navegador y escriba la dirección proporcionada por la institución.', 'En la pantalla de inicio de sesión, escriba su correo y contraseña.', 'Seleccione Iniciar sesión.', 'Al ingresar correctamente se muestra el Panel. Si la cuenta no tiene permiso para una opción, esa función no estará disponible.'])
    para(d, 'Si el sistema informa que las credenciales no coinciden, revise el correo y la contraseña. Después de varios intentos fallidos se aplica una espera temporal para proteger la cuenta.', 'Importante: ')
    heading(d, '4 Navegación y perfiles')
    table(d, ['Opción del menú', 'Uso'], [['Panel','Muestra el resumen de equipos, estado, actividad y vencimientos.'],['Inventario','Consulta, filtra, registra, edita y abre la ficha de cada equipo.'],['Préstamos','Registra entregas, devoluciones y vencimientos.'],['Mantenimiento','Consulta fallas y registra el seguimiento técnico.'],['Ubicaciones','Consulta sedes, aulas, laboratorios y otros espacios.'],['Reportes','Revisa indicadores e historial de fallas; exporta CSV o imprime.'],['Configuración','Administra opciones, catálogos, usuarios y respaldos. Solo administrador.'],['Mi perfil','Actualiza datos personales, preferencias y contraseña.']],[1.45,5.15])
    para(d, 'En teléfonos, el menú se abre con el botón de navegación superior. Las tablas se presentan como tarjetas para conservar la lectura sin desplazamiento horizontal.')
    table(d, ['Perfil', 'Alcance'], [['Administrador','Acceso completo, usuarios, configuración y eliminaciones protegidas.'],['Técnico','Inventario, préstamos, ubicaciones, mantenimiento y reportes.'],['Docente','Registro de reportes de fallas y consulta de su perfil.']], [1.25,5.35])
    heading(d, '5 Panel')
    para(d, 'El Panel presenta cuatro indicadores: equipos totales, disponibles, en préstamo y en mantenimiento. También incluye gráficos por estado y categoría, actividad reciente y próximos vencimientos. Si la institución dispone de varias sedes, el selector permite consultar los datos de una sede concreta.')
    bullets(d, ['Use las tarjetas y gráficos como resumen; no modifican registros.', 'Revise Próximos vencimientos para identificar préstamos que requieren seguimiento.', 'Abra Inventario, Préstamos o Mantenimiento desde el menú para gestionar el detalle.'])
    heading(d, '6 Inventario de equipos')
    para(d, 'El módulo Inventario lista los equipos registrados. La búsqueda localiza coincidencias por código, nombre, modelo, marca o ubicación. Los filtros por categoría, sede y estado reducen los resultados sin alterar los datos.')
    heading(d, '6.1 Registrar un equipo', 2)
    steps(d, ['Seleccione Inventario y después Registrar equipo.', 'Complete código, nombre, categoría, marca, modelo, número de serie y ubicación.', 'Seleccione responsable, sede y estado cuando corresponda.', 'Complete fecha de adquisición, precio, proveedor u observaciones si se dispone de esa información.', 'Opcionalmente cargue una fotografía JPG, PNG o WebP de hasta 3 MB.', 'Seleccione Registrar equipo y confirme el mensaje de éxito.'])
    para(d, 'El código es único. Si ya existe, el sistema no crea un segundo equipo con el mismo código.')
    heading(d, '6.2 Consultar, editar y eliminar', 2)
    steps(d, ['Busque o filtre el equipo.', 'Seleccione la fila o la opción de acciones para abrir el detalle.', 'Use Editar para actualizar la información y guarde los cambios.', 'Solo un administrador puede eliminar un registro y únicamente cuando no existan referencias que impidan la operación.'])
    heading(d, '6.3 Ficha del equipo', 2)
    para(d, 'La ficha reúne la fotografía, código, estado, ubicación, responsable, adquisición y observaciones. Sus pestañas muestran información, préstamos, mantenimientos e historial. La fotografía se entrega únicamente a usuarios autenticados con permiso.')
    heading(d, '7 Préstamos')
    para(d, 'Un préstamo conserva el estado físico del equipo y calcula su disponibilidad mientras está activo. Un equipo no puede tener dos préstamos activos al mismo tiempo ni prestarse cuando tiene una falla pendiente o está en mantenimiento.')
    heading(d, '7.1 Registrar un préstamo', 2)
    steps(d, ['Seleccione Préstamos y Nuevo préstamo.', 'Elija un equipo disponible y el usuario que lo recibe.', 'Indique fecha de préstamo, fecha prevista de devolución y observaciones si aplica.', 'Guarde el préstamo. La tabla mostrará el estado Prestado.'])
    heading(d, '7.2 Registrar una devolución', 2)
    steps(d, ['Localice el préstamo activo.', 'Abra Acciones y seleccione Registrar devolución.', 'Confirme la operación.', 'Verifique que el préstamo cambie a Devuelto y que el equipo vuelva a estar disponible.'])
    para(d, 'Un préstamo cuya fecha prevista ya pasó se identifica como Atrasado. El sistema muestra alertas al iniciar sesión cuando la opción de notificaciones está activa.')
    heading(d, '8 Reportes de fallas y mantenimiento')
    para(d, 'Los docentes pueden registrar una falla sobre un equipo. El equipo técnico utiliza el módulo Mantenimiento para ver los casos pendientes, asignar diagnóstico y cerrar el trabajo con una solución.')
    heading(d, '8.1 Reportar una falla', 2)
    steps(d, ['Abra el módulo disponible para Reportes de fallas.', 'Seleccione el equipo y describa el problema con claridad.', 'Guarde el reporte.', 'Espere la revisión del personal técnico.'])
    heading(d, '8.2 Dar seguimiento técnico', 2)
    steps(d, ['Abra Mantenimiento y filtre por Pendientes, En reparación o Completados.', 'Abra el reporte que atenderá.', 'Registre el diagnóstico y el avance.', 'Cuando el trabajo termine, escriba la solución y cierre el mantenimiento.'])
    para(d, 'El cierre exige una solución y queda como parte del historial. No se debe reabrir mantenimiento de un equipo que todavía está prestado.')
    heading(d, '9 Ubicaciones')
    para(d, 'Ubicaciones agrupa los espacios por sede y muestra cantidades de equipos, disponibles y en mantenimiento. Un administrador o técnico puede registrar y editar sedes y espacios según sus permisos.')
    steps(d, ['Abra Ubicaciones.', 'Seleccione Nueva ubicación para crear un espacio.', 'Indique sede, nombre, tipo y descripción.', 'Guarde y compruebe los conteos de la tabla.'])
    heading(d, '10 Reportes y exportación')
    para(d, 'Reportes muestra indicadores y gráficos de inventario. El historial de fallas puede filtrarse por periodo. Exportar Excel descarga un archivo CSV compatible con Excel; Exportar PDF abre una vista de impresión para guardar como PDF desde el navegador.')
    bullets(d, ['Seleccione el periodo antes de analizar el historial de fallas.', 'Abra el archivo CSV en una hoja de cálculo para ordenar o imprimir.', 'Use la ventana de impresión del navegador para guardar el informe como PDF.'])
    heading(d, '11 Configuración y perfil')
    para(d, 'Configuración es un área de administración con opciones generales, usuarios, categorías, sedes, roles de consulta, notificaciones, respaldo y apariencia. Los respaldos contienen datos privados y solo deben guardarse en ubicaciones seguras.')
    para(d, 'En Mi perfil puede actualizar nombre, teléfono, cargo, sede y densidad visual. Para cambiar la contraseña se solicita la contraseña actual y una nueva contraseña válida.')
    heading(d, '12 Mensajes frecuentes')
    table(d, ['Mensaje o situación', 'Significado', 'Acción recomendada'], [['Credenciales no válidas','El correo o contraseña no coincide.','Revise los datos o contacte al administrador.'],['El código ya existe','Otro equipo usa el mismo código.','Use un código institucional diferente.'],['Equipo no disponible','Existe préstamo o mantenimiento activo.','Registre la devolución o termine el mantenimiento.'],['Formulario caducado','La sesión o página se actualizó.','Recargue la página y vuelva a enviar el formulario.'],['Acceso denegado','El perfil no tiene permiso.','Solicite apoyo al administrador.']], [1.55,2.2,2.85])
    heading(d, '13 Preguntas frecuentes')
    table(d, ['Pregunta', 'Respuesta'], [['¿Olvidé mi contraseña?','Solicite al administrador el procedimiento institucional de restablecimiento.'],['¿Por qué un equipo no aparece como disponible?','Revise si tiene un préstamo activo o un mantenimiento sin cerrar.'],['¿Puedo borrar una ubicación?','Solo si no tiene equipos u otros registros relacionados.'],['¿El reporte genera un archivo Excel?','La descarga es CSV y puede abrirse con Excel.'],['¿Qué hago al terminar?','Cierre sesión, especialmente en equipos compartidos.']], [2.2,4.4])
    heading(d, '14 Recomendaciones de uso')
    bullets(d, ['No comparta su contraseña ni deje la sesión abierta en equipos compartidos.', 'Revise códigos, fechas y responsables antes de guardar.', 'Registre la devolución el mismo día en que el equipo regrese.', 'Describa las fallas con datos útiles para el técnico.', 'Conserve los respaldos fuera del servidor web y no los publique.'])
    heading(d, '15 Conclusión')
    para(d, 'InventIC permite que las decisiones sobre equipos se basen en un inventario, préstamos y mantenimientos registrados. El uso consistente de cada módulo mantiene la información disponible para toda la institución.')
    return d

def programmer_manual():
    d=base_doc('Manual del Programador'); cover(d, 'Manual del Programador', 'Documentación técnica para instalar, mantener y ampliar InventIC')
    heading(d, 'Control de versiones'); version_table(d)
    heading(d, '1 Introducción')
    para(d, 'InventIC - TECNO-GEST es una aplicación web institucional para inventario de equipos, préstamos, reportes de fallas y mantenimiento. Resuelve la dispersión de información sobre equipos, su ubicación, responsables, préstamos y incidencias.')
    para(d, 'Este documento está dirigido a desarrolladores, administradores técnicos y evaluadores. Describe la versión actual del código, el esquema de datos, la instalación, las reglas de negocio y los controles de seguridad implementados.')
    heading(d, '2 Descripción general y arquitectura')
    table(d, ['Capa', 'Responsabilidad', 'Componentes'], [['Cliente','Presentación responsive e interacción de baja complejidad.','HTML5, CSS, JavaScript, Bootstrap e iconos.'],['Servidor','Rutas PHP, autorización, validación y reglas de negocio.','Módulos PHP, includes, sesiones y PDO.'],['Datos','Persistencia con relaciones y restricciones.','MySQL 8, InnoDB, utf8mb4.'],['Almacenamiento','Archivos privados y control de acceso.','storage, fotografías, sesiones, respaldos y limitador.']], [1.1,2.55,2.95])
    para(d, 'Flujo de una solicitud: navegador -> Apache/PHP -> config/config.php -> autenticación y CSRF -> módulo solicitado -> helpers/consultas preparadas PDO -> MySQL -> respuesta HTML. Las escrituras críticas de préstamos y mantenimiento usan transacciones y bloqueos cuando corresponde.')
    heading(d, '3 Requisitos del sistema')
    table(d, ['Recurso', 'Mínimo recomendado'], [['Servidor','Apache 2.4 con AllowOverride habilitado.'],['PHP','PHP 8.2 o superior con PDO MySQL, mbstring, fileinfo y sesiones.'],['Base de datos','MySQL 8 con motor InnoDB y soporte utf8mb4.'],['Equipo','Procesador equivalente a Core i3, 4 GB RAM y 10 GB disponibles.'],['Cliente','Navegador moderno con JavaScript habilitado.']], [1.45,5.15])
    heading(d, '4 Tecnologías y herramientas')
    table(d, ['Tecnología', 'Uso real en el proyecto'], [['PHP 8.2','Controladores por página, formularios, sesiones y reglas de negocio.'],['PDO MySQL','Consultas preparadas, transacciones y acceso a datos.'],['MySQL 8 / InnoDB','Tablas, claves foráneas, índices, checks y columna generada.'],['Bootstrap 5.3.2','Base visual responsive cargada por CDN.'],['Bootstrap Icons 1.11.1 e Inter','Iconografía y tipografía de la interfaz.'],['CSS y JavaScript propios','Diseño InventIC, tablas móviles, filtros y comportamiento visual.'],['Git y GitHub','Control de versiones y respaldo del código fuente.'],['WAMP / Apache','Entorno local de desarrollo y pruebas.']], [1.7,4.9])
    heading(d, '5 Estructura del proyecto')
    table(d, ['Ruta', 'Función'], [['assets/','app.css y scripts de presentación responsive.'],['auth/','Inicio y cierre de sesión.'],['config/','Configuración general, conexión PDO y archivo local privado.'],['includes/','Componentes compartidos, autorización, validación y reglas de negocio.'],['modules/','Vistas y flujos por dominio: dashboard, equipos, préstamos, mantenimiento, etc.'],['tools/','Migraciones, exportación de esquema, creación de administrador y router de desarrollo.'],['tests/','Prueba de integración aislada.'],['storage/','Fotografías, respaldos, sesiones y limitador; no se publica en Git.'],['database.sql','Instalación limpia: esquema y catálogos sin cuentas ni datos privados.']], [1.55,5.05])
    heading(d, '6 Instalación desde cero')
    steps(d, ['Instale Apache, PHP 8.2+, MySQL 8 y las extensiones PHP indicadas.', 'Cree una base de datos importando database.sql. El script crea tecnogest y sus catálogos.', 'Copie el proyecto a la raíz del servidor web y permita escritura de Apache sobre storage/.', 'Defina DB_HOST, DB_PORT, DB_NAME, DB_USER y DB_PASS como variables del servidor o en config/local.php. Este archivo está excluido por .gitignore.', 'Defina temporalmente ADMIN_EMAIL y ADMIN_PASSWORD en la consola y ejecute php tools/create-admin.php. La contraseña debe tener entre 12 y 72 bytes.', 'Abra auth/login.php desde Apache y cree o use la cuenta administradora.', 'Compruebe que se puede iniciar sesión, registrar un equipo y cerrar sesión.'])
    para(d, 'Para desarrollo se puede usar php -S 127.0.0.1:8000 tools/router.php. No use el servidor integrado de PHP en producción. Configure HTTPS, permisos mínimos de base de datos y las reglas equivalentes a .htaccess si utiliza otro servidor.', 'Producción: ')
    heading(d, '7 Actualización de una instalación existente')
    para(d, 'No importe database.sql sobre una instalación con datos. Desde el esquema original ejecute primero php tools/migrate.php para convertir tablas a InnoDB y activar relaciones. Luego ejecute php tools/migrate-interface.php. Esta migración crea un respaldo privado, añade sedes, préstamos, configuración, actividad y campos ampliados sin eliminar registros.')
    para(d, 'MySQL no revierte DDL con rollback. Ejecute la migración en una ventana sin escrituras, conserve el respaldo fuera del servidor web y pruebe cualquier restauración en una base diferente antes de cambiar la conexión.')
    heading(d, '8 Base de datos')
    para(d, 'La base por defecto se denomina tecnogest, usa utf8mb4_unicode_ci e InnoDB. database.sql contiene doce tablas y catálogos mínimos; no distribuye cuentas, contraseñas, equipos de demostración ni fotografías.')
    table(d, ['Tabla', 'Propósito y claves principales'], [['roles','Catálogo de perfiles; usuarios.rol_id.'],['sedes','Sedes institucionales; ubicaciones.sede_id y usuarios.sede_id.'],['usuarios','Cuentas, hash de contraseña, perfil y preferencias. Email es único.'],['ubicaciones','Espacios físicos. Nombre único y sede_id.'],['tipos_equipo','Catálogo de categoría para equipos.'],['equipos','Activo inventariable. codigo único; referencias a tipo, ubicación y responsable.'],['estados_reporte','Catálogo de estados de incidencias.'],['reportes','Fallas reportadas por usuario sobre equipo.'],['mantenimientos','Atención técnica. Un registro por reporte mediante uq_mantenimiento_reporte.'],['prestamos','Entrega y devolución. activo_equipo garantiza un préstamo activo por equipo.'],['configuracion','Valores generales por clave.'],['actividad','Bitácora de acciones; conserva referencias opcionales.']], [1.45,5.15])
    heading(d, '8.1 Relaciones y reglas de integridad', 2)
    para(d, 'roles 1:N usuarios; sedes 1:N usuarios y ubicaciones; ubicaciones 1:N equipos; tipos_equipo 1:N equipos; equipos 1:N reportes, préstamos y actividad; usuarios 1:N reportes, mantenimientos, préstamos y actividad; reportes 1:0..1 mantenimientos. Las claves foráneas impiden eliminar entidades referenciadas. actividad usa SET NULL para conservar el evento si se elimina una referencia permitida.')
    table(d, ['Regla', 'Implementación'], [['Código de inventario único','UNIQUE equipos.codigo.'],['Correo único','UNIQUE usuarios.email.'],['Sin préstamo simultáneo','Columna STORED activo_equipo y UNIQUE uq_prestamo_activo.'],['Fechas válidas','CHECK fecha_devolucion >= fecha_prestamo y validación PHP.'],['Un mantenimiento por reporte','UNIQUE uq_mantenimiento_reporte.'],['Consultas frecuentes','Índices de estado/código, fechas de préstamo, reportes y actividad.']], [2.25,4.35])
    heading(d, '9 Componentes y módulos')
    table(d, ['Módulo', 'Archivos principales', 'Tablas'], [['Autenticación','auth/login.php, config/config.php, config/database.php','usuarios, roles'],['Panel','modules/dashboard/index.php, includes/app.php','equipos, préstamos, actividad, ubicaciones'],['Equipos','modules/equipos/, includes/equipment-form.php','equipos, tipos_equipo, ubicaciones, usuarios'],['Préstamos','modules/prestamos/, includes/loans.php','prestamos, equipos, usuarios'],['Mantenimiento','modules/mantenimiento/, includes/maintenance.php','reportes, mantenimientos, equipos'],['Ubicaciones','modules/ubicaciones/, includes/location-form.php','sedes, ubicaciones, equipos'],['Reportes','modules/reportes/, modules/estadisticas/','reportes, equipos, préstamos, mantenimientos'],['Configuración y perfil','modules/configuracion/, modules/perfil/','configuracion, usuarios, catálogos']], [1.1,3.0,2.5])
    heading(d, '10 Flujo de negocio')
    para(d, 'El inventario registra el equipo y sus relaciones. La disponibilidad se calcula: un equipo En Mantenimiento no se presta; si existe préstamo sin devuelto_en, se muestra Prestado; en los demás casos se muestra su estado físico. Los préstamos no cambian permanentemente el estado físico del equipo.')
    para(d, 'Al crear o devolver un préstamo, includes/loans.php inicia una transacción y bloquea el equipo. Antes de prestar comprueba disponibilidad, fallas pendientes y el préstamo activo. Al cerrar mantenimiento, la solución es obligatoria y el registro permanece como historial. No se abre ni reabre mantenimiento mientras el equipo está prestado.')
    heading(d, '11 Seguridad implementada')
    table(d, ['Control', 'Implementación'], [['Contraseñas','password_hash y password_verify; no se almacenan contraseñas en texto plano.'],['Sesiones','Regeneración al autenticar, HttpOnly, SameSite=Lax, Secure bajo HTTPS y caducidad por inactividad.'],['Autorización','access y admin_access verifican la sesión y rol en servidor.'],['CSRF','Token aleatorio validado en todas las solicitudes POST.'],['SQL Injection','PDO preparado en helpers y módulos; parámetros separados del SQL.'],['XSS','Salida codificada con h y htmlspecialchars.'],['Carga de archivos','JPEG/PNG/WebP, máximo 3 MB, validación de contenido, nombre aleatorio y entrega protegida.'],['Inicio de sesión','Limitador por IP: diez intentos en quince minutos.'],['Datos privados','.gitignore y .htaccess excluyen y protegen secretos, respaldos, sesiones y fotografías.'],['CSV','Neutralización de valores que podrían ser fórmulas al abrirse en hojas de cálculo.']], [1.55,5.05])
    heading(d, '12 Configuración y secretos')
    para(d, 'No incluya DB_PASS, tokens, archivos .env, config/local.php, respaldos ni contenido de storage/ en Git. El usuario MySQL de producción debe tener solo los permisos necesarios. ADMIN_EMAIL y ADMIN_PASSWORD se usan únicamente al crear una cuenta inicial y deben retirarse del entorno después de la operación.')
    para(d, 'APP_TIMEZONE y DB_TIMEZONE permiten ajustar el huso horario. Por defecto se usan America/El_Salvador y -06:00. APP_BASE_URL permite fijar la ruta base cuando la detección automática no es apropiada.')
    heading(d, '13 Pruebas y validación')
    table(d, ['Prueba', 'Resultado esperado'], [['php tests/integration.php','Ejecuta 109 comprobaciones en una base y almacenamiento temporales; no modifica la base de aplicación.'],['Migraciones','Esquema idempotente, relaciones y restauración de respaldo comprobadas.'],['Flujos','CRUD, préstamo/devolución, mantenimiento, reportes, perfil, configuración y fotografía.'],['Seguridad','CSRF, XSS, SQL Injection, permisos, sesiones, contraseña y límite de inicio de sesión.'],['Interfaz','Diez vistas recorridas en escritorio y 390 px sin desbordamiento horizontal.']], [2.2,4.4])
    para(d, 'Antes de publicar cambios ejecute php -l sobre los archivos PHP modificados, php tests/integration.php y revise git diff --check. Compruebe de forma manual el flujo afectado y revise que git status no incluya archivos privados.')
    heading(d, '14 Mantenimiento y ampliación')
    bullets(d, ['Para agregar una pantalla, reutilice page_start, page_end, navbar y los helpers de includes/app.php.', 'Para modificar el esquema, implemente una migración idempotente en includes/schema.php o tools y actualice database.sql para instalaciones limpias.', 'Mantenga validación en servidor aunque exista validación HTML o JavaScript.', 'Utilice transacciones para cambios que afecten préstamo, disponibilidad o mantenimiento.', 'Añada pruebas de integración para reglas de negocio nuevas y no elimine una restricción de base de datos sin reemplazar su protección.', 'Al restaurar una copia, copie también las fotografías privadas porque la base conserva solo sus referencias.'])
    heading(d, '15 Despliegue y recuperación')
    para(d, 'Use HTTPS, permisos de escritura solo en storage, usuario MySQL con privilegios mínimos y un directorio de respaldo ajeno al servidor web. El respaldo de Configuración es accesible únicamente a administradores autenticados por POST con CSRF. Verifique las restauraciones en una base separada. Los respaldos SQL no contienen archivos binarios de fotografías.')
    heading(d, '16 Conclusión técnica')
    para(d, 'La aplicación separa presentación, autenticación, reglas de negocio y acceso a datos con una estructura PHP directa y mantenible. El uso de InnoDB, relaciones, restricciones, consultas preparadas y pruebas de integración protege las operaciones centrales sin requerir dependencias de compilación.')
    return d

for name, builder in [('Manual de Usuario InventIC.docx', user_manual), ('Manual del Programador InventIC.docx', programmer_manual)]:
    doc = builder()
    doc.core_properties.title = name.replace('.docx','')
    doc.core_properties.author = 'InventIC'
    doc.core_properties.subject = 'Documentación del sistema InventIC'
    doc.save(OUT / name)
    print(OUT / name)
