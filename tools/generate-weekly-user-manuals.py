from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / 'docs' / 'manuales' / 'semanales'
BLUE, PALE, GRAY = '173F73', 'EFF5FC', 'D9D9D9'
MEMBERS = 'Jesús Ernesto Domínguez Reyes\nJosé Baltazar Beltrán Gómez'

WEEKS = [
    ('Semana 1', 'Acceso y panel principal',
     'Conocer el sistema, iniciar sesión y consultar el resumen institucional.',
     ['Abra la dirección institucional de InventIC en un navegador.', 'Escriba el correo y la contraseña asignados.', 'Seleccione Iniciar sesión.', 'Compruebe que se muestre el Panel y revise los indicadores de equipos.'],
     [('Panel', 'Consulta de equipos totales, disponibles, prestados y en mantenimiento.'), ('Filtro de sede', 'Actualiza la información mostrada según la sede elegida.'), ('Cerrar sesión', 'Finaliza el acceso seguro al sistema.')],
     'No comparta la contraseña. Si los datos de acceso no coinciden, verifíquelos antes de volver a intentarlo.'),
    ('Semana 2', 'Ubicaciones y consulta de inventario',
     'Ubicar los equipos por sede o espacio y encontrarlos mediante la consulta de inventario.',
     ['Ingrese a Ubicaciones desde el menú lateral.', 'Seleccione la sede o espacio que desea consultar.', 'Revise los conteos de equipos disponibles y en mantenimiento.', 'Abra Inventario y utilice búsqueda o filtros por categoría, sede o estado.'],
     [('Ubicaciones', 'Muestra sedes, aulas, laboratorios y otros espacios registrados.'), ('Inventario', 'Lista los equipos registrados con búsqueda, filtros y paginación.'), ('Ficha del equipo', 'Abre la información detallada del equipo seleccionado.')],
     'Use filtros antes de recorrer varias páginas. Los permisos determinan si puede crear, editar o solamente consultar.'),
    ('Semana 3', 'Registro y actualización de equipos',
     'Registrar equipos institucionales y mantener sus datos actualizados.',
     ['En Inventario, seleccione Registrar equipo.', 'Complete código, nombre, categoría, marca, modelo, sede y ubicación.', 'Agregue responsable, adquisición, proveedor, precio u observaciones cuando corresponda.', 'Adjunte una fotografía JPG, PNG o WebP de hasta 3 MB si es necesaria.', 'Seleccione Guardar y confirme el mensaje de registro exitoso.'],
     [('Registrar equipo', 'Crea una ficha de inventario con los datos del activo.'), ('Editar equipo', 'Actualiza datos autorizados de un equipo ya registrado.'), ('Fotografía', 'Permite conservar una imagen asociada al equipo para facilitar su identificación.')],
     'Verifique que el código sea correcto antes de guardar. No use información incompleta ni fotografías de otros equipos.'),
    ('Semana 4', 'Préstamos y devoluciones',
     'Controlar la entrega temporal de equipos y registrar su devolución.',
     ['Abra Préstamos y seleccione Nuevo préstamo.', 'Elija un equipo disponible y complete los datos del prestatario.', 'Indique las fechas de préstamo y vencimiento.', 'Guarde el préstamo y revise que aparezca en la lista.', 'Al recibir el equipo, abra el préstamo correspondiente y registre la devolución.'],
     [('Nuevo préstamo', 'Registra la salida temporal de un equipo disponible.'), ('Vencimiento', 'Permite identificar préstamos que requieren seguimiento.'), ('Devolución', 'Cierra el préstamo y conserva el historial del movimiento.')],
     'Un equipo no puede tener dos préstamos activos. Tampoco se puede prestar si tiene una falla pendiente o está en mantenimiento.'),
    ('Semana 5', 'Fallas, mantenimiento y reportes',
     'Reportar fallas, dar seguimiento técnico y consultar los indicadores del sistema.',
     ['Entre a Mantenimiento y seleccione el equipo con falla o pendiente.', 'Registre el diagnóstico y el estado del caso.', 'Cuando el trabajo termine, escriba la solución y marque el mantenimiento como completado.', 'Abra Reportes para revisar gráficos e historial de fallas.', 'Aplique el periodo deseado y exporte CSV o use la vista de impresión si lo necesita.'],
     [('Mantenimiento', 'Gestiona fallas pendientes, equipos en reparación y trabajos completados.'), ('Reportes', 'Presenta indicadores e historial filtrable de fallas.'), ('Exportar CSV', 'Descarga los datos para abrirlos en Excel u otra hoja de cálculo.')],
     'El cierre de mantenimiento requiere una solución y queda como parte del historial. Verifique los datos antes de finalizarlo.'),
    ('Semana 6', 'Perfil, configuración y operación segura',
     'Completar la operación final del sistema mediante el perfil, la configuración autorizada y buenas prácticas de uso.',
     ['Abra Mi perfil desde el menú de usuario.', 'Actualice nombre, teléfono, cargo o sede si corresponde.', 'Para cambiar la contraseña, escriba la contraseña actual y la nueva.', 'Si es administrador, abra Configuración para consultar usuarios, categorías, sedes, alertas y respaldo.', 'Cierre sesión al terminar la jornada.'],
     [('Mi perfil', 'Actualiza datos personales, preferencia visual y contraseña.'), ('Configuración', 'Área exclusiva de administración para catálogos, usuarios y respaldos.'), ('Respaldo SQL', 'Descarga una copia de la base de datos para recuperación institucional.')],
     'El respaldo contiene información institucional y debe guardarse fuera del sitio web. Solo el administrador debe usar esta función.'),
]

def font(run, size=10, bold=False, color=None):
    run.font.name = 'Aptos'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, color):
    node = OxmlElement('w:shd'); node.set(qn('w:fill'), color); cell._tc.get_or_add_tcPr().append(node)

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        cell = t.rows[0].cells[i]; cell.text = ''; shade(cell, BLUE)
        r = cell.paragraphs[0].add_run(value); font(r, 9.5, True, 'FFFFFF')
    for index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            r = p.add_run(value); font(r, 9.5)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index % 2: shade(cells[i], PALE)
    for row in t.rows:
        for cell, width in zip(row.cells, widths): cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def paragraph(doc, text, bold=False):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text); font(r, 10.5, bold)

def make_manual(num, title, focus, objective, steps, modules, note):
    d = Document(); sec = d.sections[0]
    sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
    normal = d.styles['Normal']; normal.font.name = 'Aptos'; normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos'); normal.font.size = Pt(10.5)
    for sty, size in [('Title', 24), ('Heading 1', 15), ('Heading 2', 12)]:
        s = d.styles[sty]; s.font.name = 'Aptos'; s._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos'); s.font.size = Pt(size); s.font.color.rgb = RGBColor(0, 0, 0); s.font.bold = True
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(f'InventIC | Manual de Usuario {title} | 2026'); font(r, 8, False, '666666')
    for _ in range(4): d.add_paragraph()
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run('INVENTIC IEP SAN RAFAEL'); font(r, 21, True, BLUE)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run('Sistema de Inventario de Equipos'); font(r, 13)
    d.add_paragraph()
    p = d.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f'Manual de Usuario {title}')
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(focus); font(r, 11)
    d.add_paragraph()
    table(d, ['Dato', 'Información'], [('Proyecto', 'InventIC IEP San Rafael'), ('Especialidad', 'Desarrollo de Software'), ('Integrantes', MEMBERS), ('Docente', 'Danilo Antonio Pérez Melara'), ('Periodo', f'Semana {num} de desarrollo'), ('Fecha', 'Septiembre de 2026')], [1.7, 4.9])
    d.add_page_break()
    d.add_heading('Propósito de la semana', 1); paragraph(d, objective)
    d.add_heading('Antes de comenzar', 1); paragraph(d, 'Utilice un navegador actualizado y una cuenta asignada por la institución. Las opciones visibles dependen de su perfil: administrador, técnico o docente.')
    d.add_heading('Procedimiento principal', 1)
    for step in steps: d.add_paragraph(step, style='List Number')
    d.add_heading('Opciones utilizadas', 1); table(d, ['Opción', 'Uso'], modules, [2.0, 4.6])
    d.add_heading('Resultado esperado', 1); paragraph(d, f'Al completar las acciones de {title.lower()}, la información queda disponible para los usuarios autorizados y puede consultarse desde el módulo correspondiente.')
    d.add_heading('Recomendación de uso', 1); paragraph(d, note)
    d.add_heading('Integrantes responsables', 1); table(d, ['Integrante', 'Participación documentada'], [('Jesús Ernesto Domínguez Reyes', 'Desarrollo y documentación del módulo de la semana.'), ('José Baltazar Beltrán Gómez', 'Desarrollo, pruebas y documentación del módulo de la semana.')], [2.8, 3.8])
    d.core_properties.title = f'Manual de Usuario {title} InventIC'
    d.core_properties.author = 'Jesús Ernesto Domínguez Reyes y José Baltazar Beltrán Gómez'
    d.save(OUT / f'{num:02d}_{title.replace(" ", "_").replace("y", "")}.docx')

OUT.mkdir(parents=True, exist_ok=True)
for i, week in enumerate(WEEKS, 1): make_manual(i, *week)
print(OUT)
